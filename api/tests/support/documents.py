"""Tiny knowledge base fixture files, built in the test instead of checked in."""

from __future__ import annotations


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[list[str]], *, title: str | None = None) -> bytes:
    """A valid PDF with one line of Helvetica text per entry on each page.

    An empty list for a page makes a page with no text (like a scan).
    """
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")  # filled in once the page tree exists
    pages_id = add(b"")
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids = []
    for lines in pages:
        stream = "BT /F1 11 Tf 14 TL 50 780 Td " + " ".join(
            f"({_pdf_escape(line)}) Tj T*" for line in lines
        )
        stream += " ET"
        data = stream.encode("latin-1") if lines else b""
        content = add(b"<< /Length %d >>\nstream\n" % len(data) + data + b"\nendstream")
        page_ids.append(
            add(
                (
                    f"<< /Type /Page /Parent {pages_id} 0 R "
                    f"/MediaBox [0 0 595 842] /Contents {content} 0 R "
                    f"/Resources << /Font << /F1 {font} 0 R >> >> >>"
                ).encode()
            )
        )
    kids = " ".join(f"{page} 0 R" for page in page_ids)
    objects[pages_id - 1] = (
        f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    )
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode()
    info = None
    if title:
        info = add(f"<< /Title ({_pdf_escape(title)}) >>".encode("latin-1"))

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    trailer = f"<< /Size {len(objects) + 1} /Root {catalog} 0 R"
    if info:
        trailer += f" /Info {info} 0 R"
    out += f"trailer\n{trailer} >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(path: str) -> None:
    """A Word document with a title, nested headings, a list and a table."""
    import docx

    document = docx.Document()
    document.add_paragraph("Staff Handbook", style="Title")
    document.add_heading("Opening hours", level=1)
    document.add_paragraph("We are open Monday to Friday, 9am to 5pm.")
    document.add_heading("Weekends", level=2)
    document.add_paragraph("On Saturdays we open from 10am to 2pm.")
    document.add_paragraph("Sundays are closed", style="List Bullet")
    document.add_paragraph("Bank holidays are closed", style="List Bullet")
    document.add_heading("Prices", level=1)
    table = document.add_table(rows=3, cols=2)
    for row, (item, price) in zip(
        table.rows, [("Item", "Price"), ("Coffee", "3.00"), ("Tea", "2.50")]
    ):
        row.cells[0].text = item
        row.cells[1].text = price
    document.add_paragraph("Prices include VAT.")
    document.save(path)
